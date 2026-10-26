"""Text extraction. Returns a list of (page_number | None, text) segments."""
from pathlib import Path

Segment = tuple[int | None, str]


class ExtractionError(Exception):
    pass


def _pdf(path: Path) -> list[Segment]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ExtractionError("This PDF is password-protected.")
        segments = []
        for i, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                segments.append((i, text))
    except ExtractionError:
        raise
    except Exception as e:  # corrupt file etc.
        raise ExtractionError(f"Could not read the PDF: {e}")
    return segments


def _docx(path: Path) -> list[Segment]:
    from docx import Document

    try:
        doc = Document(str(path))
    except Exception as e:
        raise ExtractionError(f"Could not read the Word file: {e}")

    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:  # course outlines are table-heavy (weekly plan, grading)
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = c.text.strip().replace("\n", " ")
                if t and (not cells or cells[-1] != t):  # skip merged-cell repeats
                    cells.append(t)
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return [(None, text)] if text else []


def _text(path: Path) -> list[Segment]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc).strip()
            return [(None, text)] if text else []
        except UnicodeError:
            continue
    return []


def extract_text(path: Path) -> list[Segment]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        segments = _pdf(path)
    elif ext == ".docx":
        segments = _docx(path)
    elif ext in (".txt", ".md"):
        segments = _text(path)
    else:
        raise ExtractionError(f"Unsupported file type: {ext}")

    if not segments:
        raise ExtractionError(
            "No readable text found. If this is a scanned PDF, run OCR on it first."
        )
    return segments
